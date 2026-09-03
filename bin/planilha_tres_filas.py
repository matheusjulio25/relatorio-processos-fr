#!/usr/bin/env python3
"""Planilha das 3 filas de trabalho: DESPACHAR, LER LAUDO, COBRAR LAUDO —
com a providência concreta de cada processo."""
import json, collections
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND="7F8187"; SUB="5F6066"; VERDE="1E7E34"; VERM="C0392B"; LARANJA="B7791F"; CINZA="777777"
HOJE=datetime(2026,9,3)
K={r['c']:r for r in json.load(open('mapas/kpi_0309.json',encoding='utf-8'))['rows']}
M={x['cnj']:x for x in json.load(open('mapas/maturidade_0309_final.json',encoding='utf-8')) if x['rodada']=='03/09/2026'}
M2=json.load(open('mapas/mapa_102.json',encoding='utf-8'))
PER=json.load(open('mapas/aba_pericia.json',encoding='utf-8'))
G=json.load(open('mapas/grupos_acao.json',encoding='utf-8'))
CNPJ="37.339.869/0001-75"
def d2(s):
    try: return datetime.strptime(s,'%d/%m/%Y')
    except: return None
def uteis(a,b):
    n=0; d=a
    while d<b:
        d+=timedelta(days=1)
        if d.weekday()<5: n+=1
    return n

# ---------- providências por fila ----------
def prov_despachar(c):
    r=K[c]; m=M.get(c,{}); base=(M2.get(c) or ['',''])[1]
    t187=r['t']==1; conc=m.get('conclusos_em'); soc=m.get('pericia_social')
    pas=["REQUERER O JULGAMENTO de procedência — instrução encerrada."]
    if conc:
        dd=d2(conc); dias=(HOJE-dd).days if dd else 0
        pas.append(f"Autos conclusos desde {conc} ({dias} dias) — pedir julgamento imediato.")
    if t187:
        pas.append("Invocar o TEMA 187/TNU: o indeferimento foi pelo critério deficiência, "
                   "a renda é incontroversa e a perícia social é dispensável. Conferir antes as 4 condições, "
                   "em especial o prazo de 2 anos entre o indeferimento e a demanda.")
    elif soc=='designada_pendente':
        pas.append("Perícia social ainda pendente: requerer a dispensa ou, subsidiariamente, "
                   "a designação imediata com prazo, para não perpetuar a paralisação.")
    if not m.get('manif_laudo_em'):
        pas.append("Não há manifestação nossa sobre o laudo — manifestar concordância no mesmo ato.")
    pas.append("Pedir DIB na DER.")
    pas.append(f"Informar a retenção de 30% e pedir destaque/RPV em nome do CNPJ {CNPJ}.")
    return " ".join(pas), base

def prov_ler(c):
    r=K[c]; m=M.get(c,{}); base=(M2.get(c) or ['',''])[1]
    conc=m.get('conclusos_em'); dt=m.get('laudo_medico_em') or m.get('laudo_social_em')
    pas=[f"LER o laudo de {dt} no PJe e classificar entre despachar e impugnar." if dt
         else "LER o laudo no PJe e classificar entre despachar e impugnar."]
    pas.append("No documento, procurar: (a) a resposta ao quesito do IMPEDIMENTO DE LONGO PRAZO "
               "(Sim/Não); (b) o nível de suporte, quando TEA; (c) o escore da Matriz IF-Br; "
               "(d) se o perito enfrentou as BARREIRAS e os quesitos do autor.")
    if conc: pas.append(f"PRIORIDADE — autos conclusos desde {conc}: o juiz pode sentenciar a qualquer momento.")
    if r['t']==1: pas.append("Tema 187 aplicável: qualquer que seja o resultado, assentar a dispensa da social.")
    pas.append("Atenção: o visualizador do PJe costuma abrir a CAPA do laudo; o conteúdo está no "
               "documento de id VIZINHO (o anexo).")
    return " ".join(pas), base

def prov_cobrar(c):
    m=M.get(c,{}); v=PER.get(c,{})
    ps=[p for p in v.get('pericias',[]) if d2(p['data']) and d2(p['data'])<HOJE and 'Cancel' not in p['situacao']]
    ps=sorted(ps,key=lambda p:d2(p['data']),reverse=True)
    p0=ps[0] if ps else {}
    dd=d2(p0.get('data','')) if p0 else None
    du=uteis(dd,HOJE) if dd else 0
    perito=(p0.get('perito','') or '').split('-')[0].strip()
    pas=[f"COBRAR a juntada do laudo — perícia realizada em {p0.get('data','?')}"
         f"{' pelo(a) perito(a) '+perito if perito else ''}, "
         f"há {du} dias úteis, sem laudo nos autos."]
    pas.append("Peticionar requerendo a intimação do perito para juntada em prazo certo, "
               "sob pena de substituição e de comunicação ao juízo para nova nomeação (art. 468, II, CPC).")
    pas.append("Registrar que a paralisação decorre exclusivamente da ausência do laudo, "
               "para afastar qualquer alegação de desídia da parte.")
    if m.get('conclusos_em'):
        pas.append(f"ATENÇÃO: autos conclusos desde {m['conclusos_em']} sem o laudo — alertar o juízo.")
    return " ".join(pas), f"Perícia em {p0.get('data','?')} · situação '{p0.get('situacao','?')}' · {du} dias úteis"

FILAS=[("1. Despachar","Despachar",prov_despachar,VERDE,"5 dias",
        "Laudo pericial FAVORÁVEL lido no inteiro teor. Instrução encerrada — requerer o julgamento."),
       ("2. Ler laudo","Ler laudo",prov_ler,LARANJA,"5 dias",
        "Laudo entregue nos autos mas a conclusão não foi recuperável pela extração automática. Ler no PJe."),
       ("3. Cobrar laudo","Cobrar laudo",prov_cobrar,VERM,"5 dias",
        "Perícia realizada há mais de 15 dias úteis e o laudo não foi juntado.")]

wb=Workbook(); wb.remove(wb.active)
thin=Side(style="thin",color="D8DCE1")
for nome,chave,fn,cor,prazo,desc in FILAS:
    ws=wb.create_sheet(nome)
    ws["A1"]=f"Fernandes & Rêgo Advogados Associados — {nome.split('. ')[1].upper()}"
    ws["A1"].font=Font(bold=True,size=14,color=BRAND,name="Palatino Linotype")
    ws["A2"]=f"{desc}  ·  Rodada 03/09/2026 · prazo sugerido: {prazo}"
    ws["A2"].font=Font(italic=True,size=9,color=SUB,name="Segoe UI")
    ws.append([])
    COLS=["#","Processo (CNJ)","Cliente","Vara","T187","Conclusos desde",
          "Base — o que o laudo diz","Providência a tomar"]
    ws.append(COLS)
    for c in ws[4]:
        c.font=Font(bold=True,color="FFFFFF",name="Segoe UI",size=10)
        c.fill=PatternFill("solid",fgColor=BRAND)
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    itens=G.get(chave,[])
    def ordem(c):
        m=M.get(c,{}); d=d2(m.get('conclusos_em','')) 
        return (0 if d else 1, d or datetime(2030,1,1), c)
    for i,c in enumerate(sorted(itens,key=ordem),1):
        r=K[c]; m=M.get(c,{})
        prov,base=fn(c)
        ws.append([i,c,r.get('n','') or m.get('cliente','') or '—',
                   m.get('vara','') or r.get('vara',''),
                   "SIM" if r['t']==1 else "",
                   m.get('conclusos_em','') or "—", base or "—", prov])
        row=ws[ws.max_row]
        for cel in row:
            cel.font=Font(name="Segoe UI",size=9)
            cel.alignment=Alignment(vertical="top",wrap_text=True)
            cel.border=Border(bottom=thin)
        row[1].font=Font(name="Segoe UI",size=9,bold=True)
        if r['t']==1: row[4].font=Font(name="Segoe UI",size=9,bold=True,color=VERDE)
        if m.get('conclusos_em'): row[5].font=Font(name="Segoe UI",size=9,bold=True,color=VERM)
        row[7].font=Font(name="Segoe UI",size=9,color=cor)
    for j,w in enumerate([5,23,30,20,7,15,62,86],1):
        ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes="A5"
    ws.auto_filter.ref=f"A4:{get_column_letter(len(COLS))}{ws.max_row}"
    ws.row_dimensions[1].height=20

# capa
cap=wb.create_sheet("Resumo",0)
cap["A1"]="Três filas de trabalho — 03/09/2026"
cap["A1"].font=Font(bold=True,size=16,color=BRAND,name="Palatino Linotype")
cap["A2"]="Providências geradas a partir da leitura do inteiro teor dos laudos no PJe"
cap["A2"].font=Font(italic=True,size=10,color=SUB,name="Segoe UI")
cap.append([]); cap.append(["Fila","Qtd","Conclusos","Tema 187","Prazo","O que é"])
for c in cap[4]:
    c.font=Font(bold=True,color="FFFFFF",name="Segoe UI"); c.fill=PatternFill("solid",fgColor=BRAND)
    c.alignment=Alignment(horizontal="center",wrap_text=True)
for nome,chave,fn,cor,prazo,desc in FILAS:
    it=G.get(chave,[])
    cap.append([nome,len(it),
                sum(1 for c in it if M.get(c,{}).get('conclusos_em')),
                sum(1 for c in it if K[c]['t']==1), prazo, desc])
    for c in cap[cap.max_row]:
        c.font=Font(name="Segoe UI",size=10); c.alignment=Alignment(wrap_text=True,vertical="top")
    cap[cap.max_row][0].font=Font(name="Segoe UI",size=10,bold=True,color=cor)
cap.append([])
cap.append(["TOTAL",sum(len(G.get(k,[])) for _,k,_,_,_,_ in FILAS),"","","",""])
for c in cap[cap.max_row]: c.font=Font(bold=True,name="Segoe UI")
for j,w in enumerate([18,8,12,11,12,78],1): cap.column_dimensions[get_column_letter(j)].width=w

out="relatorios/Filas_Despachar_Ler_Cobrar_03set2026.xlsx"; wb.save(out)
print("->",out)
for nome,chave,_,_,_,_ in FILAS: print(f"   {nome}: {len(G.get(chave,[]))}")
