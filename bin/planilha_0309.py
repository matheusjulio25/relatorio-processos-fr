#!/usr/bin/env python3
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
BRAND="7F8187"; SUB="5F6066"; VERM="C0392B"; VERDE="1E7E34"; AMB="B7791F"
M={r['cnj']:r for r in json.load(open('mapas/maturidade_0309_final.json',encoding='utf-8'))}
F=json.load(open('mapas/filas_0309.json',encoding='utf-8'))
V=json.load(open('mapas/laudo_veredito.json',encoding='utf-8'))
COLS=["Processo (CNJ)","Cliente","Vara","Benefício","Estado","P. médica","P. social",
      "Laudo","Fonte do laudo","Motivo indef.","Tema 187","Laudo em","Manif. laudo",
      "Conclusos em","Últ. mov.","Novo no acervo","Próxima diligência"]
ABAS=[("1b URGENTE laudo adverso","1b_URGENTE_conclusos_laudo_adverso",
       "CONCLUSOS com laudo DESFAVORÁVEL — improcedência iminente. IMPUGNAR o laudo antes de qualquer pedido de julgamento."),
      ("A laudo novo sem manif","A_laudo_novo_sem_manifestacao",
       "Laudo juntado após 10/07 SEM manifestação do escritório — prazo correndo."),
      ("1 Despachar favoravel","1_DESPACHAR_laudo_favoravel",
       "Ativo + laudo judicial FAVORÁVEL — requerer julgamento de procedência."),
      ("1c Tema 187","1c_TEMA187",
       "Ativo + indeferimento por deficiência + social não realizada — pedir dispensa da social (Tema 187/TNU)."),
      ("2 Laudo nao lido","2_laudo_nao_lido","Laudo existe mas o resultado NÃO foi possível apurar — ler antes de decidir."),
      ("3 Aguardando pericia","3_aguardando_pericia","Ativo sem nenhum laudo — cobrar designação + protocolo antifalta D-3/D-1."),
      ("4 Em recurso","4_em_recurso","Na Turma Recursal — nada a despachar em 1º grau."),
      ("5 Cumprimento","5_cumprimento","Fase de cumprimento — cálculo/RPV."),
      ("6 Sentenciado","6_sentenciado","Com sentença — conferir prazo recursal."),
      ("7 Arquivado","7_arquivado","Baixa definitiva — nada a fazer, salvo reajuizamento.")]
def dilig(r,fila):
    if fila.startswith("1b"): return "IMPUGNAR o laudo (autos conclusos) + assentar dispensa da social pelo Tema 187, se aplicável"
    if fila.startswith("A"):  return "Manifestar sobre o laudo juntado — prazo correndo"
    if fila.startswith("1_"): return "Requerer JULGAMENTO de procedência (instrução encerrada)"
    if fila.startswith("1c"): return "Requerer dispensa da perícia social (Tema 187/TNU) + julgamento"
    if fila.startswith("2"):  return "LER o laudo e decidir entre despachar e impugnar"
    if fila.startswith("3"):  return "Cobrar designação da perícia + protocolo antifalta D-3/D-1"
    if fila.startswith("4"):  return "Atuar na Turma Recursal (memorial)"
    if fila.startswith("5"):  return "Homologação de cálculo / RPV"
    if fila.startswith("6"):  return "Conferir prazo recursal e recorrer, se cabível"
    return "Nada — arquivado"
wb=Workbook(); wb.remove(wb.active)
for nome,chave,desc in ABAS:
    ws=wb.create_sheet(nome[:31])
    ws["A1"]="Fernandes & Rêgo Advogados Associados — Maturidade de Julgamento"
    ws["A1"].font=Font(bold=True,size=13,color=BRAND,name="Palatino Linotype")
    ws["A2"]=f"{desc}  ·  Rodada 03/09/2026 · fonte: PJe (leitura de timeline hoje)"
    ws["A2"].font=Font(italic=True,size=9,color=SUB,name="Segoe UI")
    ws.append([]); ws.append(COLS)
    hdr=ws[4]
    for c in hdr:
        c.font=Font(bold=True,color="FFFFFF",name="Segoe UI",size=9)
        c.fill=PatternFill("solid",fgColor=BRAND); c.alignment=Alignment(wrap_text=True,vertical="center")
    for cnj in F[chave]:
        r=M[cnj]
        ws.append([r['cnj'],r['cliente'],r['vara'],r['classe_beneficio'],r['estado_processo'],
                   r['pericia_medica'],r['pericia_social'],r['resultado_laudo'],r['laudo_fonte'],
                   r['motivo_indeferimento'],
                   {True:"SIM",False:"não",None:"n/d"}[r['dispensa_social_tema187']],
                   r['laudo_medico_em'] or r['laudo_social_em'],r['manif_laudo_em'],
                   r['conclusos_em'],r['ultima_data'],"NOVO" if r['novo_no_acervo'] else "",
                   dilig(r,chave)])
        row=ws[ws.max_row]
        for c in row: c.font=Font(name="Segoe UI",size=9); c.alignment=Alignment(vertical="top",wrap_text=False)
        cor={'desfavoravel':VERM,'favoravel':VERDE}.get(r['resultado_laudo'],AMB)
        row[7].font=Font(name="Segoe UI",size=9,bold=True,color=cor)
        if r['dispensa_social_tema187']: row[10].font=Font(name="Segoe UI",size=9,bold=True,color=VERDE)
    larg=[24,32,20,17,20,17,18,14,13,18,9,12,12,12,12,7,62]
    for i,w in enumerate(larg,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:{get_column_letter(len(COLS))}{ws.max_row}"
# capa
cap=wb.create_sheet("Panorama",0)
cap["A1"]="Maturidade de Julgamento — Fernandes & Rêgo"
cap["A1"].font=Font(bold=True,size=16,color=BRAND,name="Palatino Linotype")
cap["A2"]="Rodada 03/09/2026 · 777 processos no acervo · 423 vivos lidos no PJe hoje"
cap["A2"].font=Font(italic=True,size=10,color=SUB,name="Segoe UI")
cap.append([]); cap.append(["Fila","Qtd","O que fazer"])
for c in cap[4]:
    c.font=Font(bold=True,color="FFFFFF",name="Segoe UI"); c.fill=PatternFill("solid",fgColor=BRAND)
for nome,chave,desc in ABAS:
    cap.append([nome,len(F[chave]),desc])
    for c in cap[cap.max_row]: c.font=Font(name="Segoe UI",size=10); c.alignment=Alignment(wrap_text=True,vertical="top")
for i,w in enumerate([30,8,95],1): cap.column_dimensions[get_column_letter(i)].width=w
out="relatorios/Maturidade_Julgamento_03set2026.xlsx"; wb.save(out); print("->",out)
