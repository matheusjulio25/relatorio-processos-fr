#!/usr/bin/env python3
"""Gera um .docx com a identidade visual COMPLETA do escritório Fernandes & Rêgo:
capa com logo, tipografia Palatino Linotype (títulos) / Segoe UI (corpo), cor da
marca #7F8187, tabelas estilizadas (cabeçalho na cor da marca + zebra), cabeçalho
com logo, rodapé institucional (doc · data · versão · página), A4 e margens.

Uso: python3 bin/relatorio_fr_docx.py <base.docx pandoc> <saida.docx> <logo.png> "<titulo>" "<subtitulo>" "<data>"
"""
import sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE, OUT, LOGO = sys.argv[1], sys.argv[2], sys.argv[3]
TITULO = sys.argv[4] if len(sys.argv) > 4 else "Relatório"
SUBTITULO = sys.argv[5] if len(sys.argv) > 5 else ""
DATA = sys.argv[6] if len(sys.argv) > 6 else ""

MARCA = "7F8187"; MARCA_ESC = "5F6066"; CORPO = "242424"; RISCO = "C00000"; CLARO = "F1F0F2"
SERIF = "Palatino Linotype"; SANS = "Segoe UI"

doc = Document(BASE)

def set_fonte(style, nome, size=None, cor=None, bold=None):
    style.font.name = nome
    if size is not None: style.font.size = Pt(size)
    if cor is not None: style.font.color.rgb = RGBColor.from_string(cor)
    if bold is not None: style.font.bold = bold
    rpr = style.element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(a), nome)

# --- estilos ---
nomes = {s.name for s in doc.styles}
set_fonte(doc.styles["Normal"], SANS, 10.5, CORPO)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
for nm, fonte, sz, cor in [("Title", SERIF, 26, MARCA), ("Heading 1", SERIF, 17, MARCA),
                            ("Heading 2", SERIF, 13.5, MARCA_ESC), ("Heading 3", SANS, 11.5, MARCA_ESC),
                            ("Heading 4", SANS, 10.5, MARCA_ESC)]:
    if nm in nomes:
        set_fonte(doc.styles[nm], fonte, sz, cor, bold=True)
        doc.styles[nm].paragraph_format.space_before = Pt(12 if nm.startswith("Heading 1") else 8)
        doc.styles[nm].paragraph_format.space_after = Pt(4)

def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)

def cell_text(cell, color=None, bold=None):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = SANS; r.font.size = Pt(9.5)
            if color: r.font.color.rgb = RGBColor.from_string(color)
            if bold is not None: r.font.bold = bold

# --- tabelas: cabeçalho na cor da marca + zebra ---
for tbl in doc.tables:
    tbl.alignment = 1
    for ci, cell in enumerate(tbl.rows[0].cells):
        shade(cell, MARCA); cell_text(cell, "FFFFFF", True)
    for ri, row in enumerate(tbl.rows[1:], start=1):
        for cell in row.cells:
            if ri % 2 == 0: shade(cell, CLARO)
            cell_text(cell)

# --- capa (inserida no topo) ---
covers = []
pL = doc.add_paragraph(); pL.alignment = WD_ALIGN_PARAGRAPH.CENTER
pL.add_run().add_picture(LOGO, width=Cm(9)); pL.paragraph_format.space_before = Pt(90); covers.append(pL._p)
def cap(txt, fonte, sz, cor, bold=True, sp_before=0, sp_after=6, italic=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.font.name = fonte; r.font.size = Pt(sz)
    r.font.color.rgb = RGBColor.from_string(cor); r.font.bold = bold; r.font.italic = italic
    p.paragraph_format.space_before = Pt(sp_before); p.paragraph_format.space_after = Pt(sp_after)
    covers.append(p._p); return p
cap(TITULO, SERIF, 24, MARCA, sp_before=60, sp_after=10)
if SUBTITULO: cap(SUBTITULO, SANS, 12, MARCA_ESC, bold=False, sp_after=40)
if DATA: cap(DATA, SANS, 10.5, MARCA_ESC, bold=False, sp_before=30, sp_after=4)
cap("Documento interno de inteligência jurídica · Confidencial", SANS, 9, MARCA, bold=False, italic=True, sp_after=0)
# quebra de página ao fim da capa
pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE); covers.append(pb._p)

body = doc.element.body
anchor = body[0]
for el in covers: body.remove(el)
for el in covers: anchor.addprevious(el)

# --- página A4 + margens + cabeçalho/rodapé ---
sec = doc.sections[0]
sec.page_height = Twips(16838); sec.page_width = Twips(11906)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"): setattr(sec, m, Twips(1440))

hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run().add_picture(LOGO, width=Cm(4))

fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run(f"Fernandes & Rêgo Advogados Associados  ·  {TITULO}  ·  {DATA}  ·  Versão 1.0  ·  ")
fr.font.name = SANS; fr.font.size = Pt(8); fr.font.color.rgb = RGBColor.from_string(MARCA)
# nº de página
run = fp.add_run(); run.font.name = SANS; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(MARCA)
fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

doc.save(OUT)
print("OK:", OUT)
