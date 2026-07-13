#!/usr/bin/env python3
"""Planilha da Controladoria a partir das DECISÕES DA INTELIGÊNCIA (não da contagem).

Lê o journal do workflow (todas as análises) + a pauta (nome do periciado/órgão),
e monta o .xlsx com abas por PRIORIDADE DE AÇÃO. Copia p/ o Drive Controladoria.

Uso: python3 bin/planilha_inteligencia.py [journal.jsonl]
"""
from __future__ import annotations
import glob
import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
PAUTA = BASE / "mapas" / "pericias_por_processo.json"
OUT = BASE / "relatorios" / "Maturidade_Julgamento_IA.xlsx"
DRIVE = Path("/Users/matheusjuliorego/Library/CloudStorage/"
             "GoogleDrive-matheus@julioregoadv.com/Drives compartilhados/"
             "Controladoria - Planilhas de Controle")
MARCA = "7F8187"


def _deep(o):
    if isinstance(o, dict):
        if "estado_processo" in o:
            return o
        for v in o.values():
            r = _deep(v)
            if r:
                return r
    if isinstance(o, str):
        try:
            return _deep(json.loads(o))
        except Exception:
            return None
    return None


def carregar_analises(journal):
    by = {}
    if str(journal).endswith(".json"):
        for c in json.load(open(journal, encoding="utf-8")):
            if c.get("cnj"):
                by[c["cnj"]] = c
        return by
    for ln in open(journal, encoding="utf-8"):
        try:
            d = json.loads(ln)
        except Exception:
            continue
        if d.get("type") != "result":
            continue
        c = _deep(d)
        if c and c.get("cnj"):
            by[c["cnj"]] = c  # última vence
    return by


# aba -> (título/ação, filtro)
def bucket(a):
    est = a.get("estado_processo")
    if a.get("despachar_agora") is True:
        lr = a.get("resultado_laudo")
        if lr == "favoravel":
            return "1. DESPACHAR - laudo favoravel"
        if lr in ("desfavoravel", "inconclusivo", "nao_verificado"):
            return "1b. REVISAR antes - laudo adverso"
        return "1c. DESPACHAR - sem laudo (Tema 187)"
    if est == "ativo_sem_sentenca":
        # ativo mas não pronto: aguardando perícia/laudo/impugnação
        if a.get("pericia_social") in ("designada_pendente", "nao_ha") and a.get("classe_beneficio") == "bpc_deficiencia" and not a.get("dispensa_social_tema187"):
            return "2. Aguardando perícia social"
        return "3. Ativo — outra diligência"
    if est == "em_recurso":
        return "4. Em recurso (Turma)"
    if est == "cumprimento_sentenca":
        return "5. Cumprimento de sentença"
    if est == "sentenciado":
        rs = a.get("resultado_sentenca")
        if rs in ("procedente", "parcialmente_procedente"):
            return "6a. Sentenciado - procedente"
        if rs == "improcedente":
            return "6b. Sentenciado - improcedente (recorrer?)"
        return "6c. Sentenciado - outro"
    if est in ("arquivado_extinto",):
        return "7. Arquivado/extinto"
    return "8. Outros/revisar"


def extrai_vara(a):
    t = " ".join(str(a.get(k, "") or "") for k in ("juiz", "fundamento", "proxima_diligencia"))
    m = re.search(r"(\d{1,2})[ªa]?\s*Vara", t)
    return f"{m.group(1)}ª Vara" if m else ""


COLS = [
    ("cnj", "Processo (CNJ)", 22), ("cliente", "Cliente", 24),
    ("vara", "Vara", 11), ("classe_beneficio", "Benefício", 15),
    ("estado_processo", "Estado", 18), ("pericia_medica", "P. médica", 12),
    ("pericia_social", "P. social", 12), ("resultado_laudo", "Laudo", 14),
    ("motivo_indeferimento", "Motivo indef.", 14), ("dispensa_social_tema187", "Tema 187", 9),
    ("resultado_sentenca", "Sentença", 16), ("juiz", "Juiz", 22), ("perito", "Perito", 22),
    ("proxima_diligencia", "Próxima diligência", 60), ("texto_peticao", "Minuta", 70),
]


def main():
    journal = sys.argv[1] if len(sys.argv) > 1 else sorted(
        glob.glob(str(BASE / ".claude" ) ))  # fallback overwritten below
    if len(sys.argv) <= 1:
        cands = glob.glob("/Users/matheusjuliorego/.claude/projects/*/subagents/workflows/*/journal.jsonl")
        cands.sort(key=lambda p: Path(p).stat().st_mtime)
        journal = cands[-1] if cands else ""
    by = carregar_analises(journal) if journal and Path(journal).exists() else {}

    pauta = json.loads(PAUTA.read_text(encoding="utf-8"))

    def cliente(cnj):
        return (pauta.get(cnj, {}) or {}).get("cliente", "")

    def orgao(cnj):
        return (pauta.get(cnj, {}) or {}).get("orgao", "")

    rows = []
    for cnj, a in by.items():
        a = dict(a)
        a["cliente"] = cliente(cnj)
        a["orgao"] = orgao(cnj)
        a["vara"] = extrai_vara(a)
        a["_bucket"] = bucket(a)
        rows.append(a)

    def escrever_aba(nome, sub):
        wsx = wb.create_sheet(re.sub(r"[/\\?*:\[\]]", "-", nome)[:31])
        for j, (_, t, w) in enumerate(COLS, 1):
            cc = wsx.cell(row=1, column=j, value=t)
            wsx.column_dimensions[get_column_letter(j)].width = w
            cc.fill = PatternFill("solid", fgColor=MARCA)
            cc.font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
            cc.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        wsx.freeze_panes = "A2"
        for i, r in enumerate(sub, 2):
            for j, (key, _, _) in enumerate(COLS, 1):
                v = r.get(key, "")
                if isinstance(v, bool):
                    v = "Sim" if v else "Não"
                c = wsx.cell(row=i, column=j, value=v)
                c.font = Font(name="Segoe UI", size=10)
                c.alignment = Alignment(vertical="top", wrap_text=(key in ("proxima_diligencia", "texto_peticao")))
        wsx.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(len(sub)+1,2)}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    hoje = date.today().strftime("%d/%m/%Y")
    from collections import Counter
    cnt = Counter(r["_bucket"] for r in rows)
    linhas = [["Maturidade p/ Julgamento — DECISÃO DA INTELIGÊNCIA (Fernandes & Rêgo)", ""],
              ["Gerado em", hoje], ["Processos analisados", len(rows)], ["", ""]]
    for b in sorted(cnt):
        linhas.append([b, cnt[b]])
    for i, (a1, b1) in enumerate(linhas, 1):
        ws.cell(row=i, column=1, value=a1).font = Font(bold=(i == 1), name="Segoe UI", size=12 if i == 1 else 10)
        ws.cell(row=i, column=2, value=b1).font = Font(name="Segoe UI", size=10)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 12

    # aba consolidada: TODOS os processos COM SENTENÇA (para visualizar e despachar)
    RS_ORD = {"procedente": 0, "parcialmente_procedente": 1, "homologacao": 2,
              "improcedente": 3, "extincao_sem_merito": 4}
    com_sent = [r for r in rows if r.get("resultado_sentenca") not in (None, "sem_sentenca", "nao_verificado")]
    com_sent.sort(key=lambda r: (RS_ORD.get(r.get("resultado_sentenca"), 9), r.get("vara", "")))
    escrever_aba("0. COM SENTENCA (todos)", com_sent)

    for b in sorted(set(r["_bucket"] for r in rows)):
        escrever_aba(b, [r for r in rows if r["_bucket"] == b])

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"[xlsx] {OUT} | {len(rows)} processos | buckets: {dict(cnt)}")
    if DRIVE.exists():
        shutil.copy2(OUT, DRIVE / OUT.name)
        print(f"[drive] copiado p/ {DRIVE / OUT.name}")
    else:
        print(f"[drive] AVISO: pasta não encontrada")


if __name__ == "__main__":
    main()
